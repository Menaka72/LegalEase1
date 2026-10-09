from io import BytesIO
from fastapi.responses import StreamingResponse
from docx import Document
from docx.shared import Pt
from fpdf import FPDF


def format_docx(text, document_type):
    doc = Document()
    title = doc.add_heading(document_type, level=0)
    title.alignment = 1

    for paragraph in text.split("\n"):
        if paragraph.strip():
            p = doc.add_paragraph(paragraph.strip())
            for run in p.runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(12)

    output = BytesIO()
    doc.save(output)
    output.seek(0)

    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": "attachment; filename=legal_document.docx"}
    )


def format_pdf(text, document_type):
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()

    pdf.set_font("Times", "B", 16)
    pdf.cell(0, 10, document_type, ln=True, align="C")
    pdf.ln(5)

    pdf.set_font("Times", size=12)
    for paragraph in text.split("\n"):
        if paragraph.strip():
            pdf.multi_cell(0, 7, paragraph.strip())
            pdf.ln(2)

    data = bytes(pdf.output())
    return StreamingResponse(
        BytesIO(data),
        media_type="application/pdf",
        headers={"Content-Disposition": "attachment; filename=legal_document.pdf"}
    )
